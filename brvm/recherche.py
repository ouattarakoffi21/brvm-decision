"""Module 9 : rapport de recherche par valeur, adapté au profil de l'investisseur.

Chaque indicateur est calculé à partir de données publiées (comptes annuels
repris des fiches société Sikafinance, cours et dividendes validés par le
module qualité). Les notes (soutenabilité, solidité, risque) suivent des
règles fixes décrites dans la page Méthodologie : aucune appréciation
subjective, aucun chiffre estimé à la main.

Conventions :
- BNPA = résultat net / nombre de titres actuel. Les cours utilisés sont
  ajustés des divisions de nominal, donc comparables au BNPA.
- La fourchette à 12 mois n'est PAS une prévision : elle applique au cours
  actuel les variations sur 12 mois que le titre a réellement connues sur
  10 ans au plus (périodes de hausse et de baisse) (2 cas sur 10 en dessous du bas, 2 cas sur 10 au-dessus du haut).
"""
from __future__ import annotations

import numpy as np
import pandas as pd

from .config import chemin

# Pondérations du score de profil (sur 100), par tolérance au risque.
# score = score de l'outil ; securite = 10 - note de risque ;
# soutenabilite = note de dividende ; croissance = rang du TCAM du CA ;
# momentum = rang du momentum 6 mois.
PONDERATIONS_PROFIL = {
    "prudent": {"score": 40, "securite": 30, "soutenabilite": 30},
    "équilibré": {"score": 50, "securite": 20, "soutenabilite": 15, "croissance": 15},
    "dynamique": {"score": 40, "croissance": 30, "momentum": 30},
}
# Note de risque maximale acceptée dans le top 10 selon la tolérance.
RISQUE_MAX = {"prudent": 4.5, "équilibré": 6.5, "dynamique": 10.0}
PER_PLAUSIBLE = (1.0, 50.0)  # au-delà, le bénéfice est trop faible pour que le PER ait un sens


def charger_comptes(cfg: dict) -> pd.DataFrame:
    """Comptes annuels publiés (CA, résultat net, dividende, nombre de titres)."""
    p = chemin(cfg, "dossier_archive") / "fondamentaux_societes.csv"
    if not p.exists():
        return pd.DataFrame(columns=["ticker", "exercice"])
    df = pd.read_csv(p)
    df["bnpa"] = df["resultat_net_mfcfa"] * 1e6 / df["nombre_titres"]
    return df.sort_values(["ticker", "exercice"])


def _tcam(serie: pd.Series) -> float:
    """Taux de croissance annuel moyen entre le premier et le dernier exercice positifs."""
    s = serie.dropna()
    if len(s) < 2 or s.iloc[0] <= 0 or s.iloc[-1] <= 0:
        return np.nan
    return (s.iloc[-1] / s.iloc[0]) ** (1 / (len(s) - 1)) - 1


def _rang(s: pd.Series, croissant: bool = True) -> pd.Series:
    """Rang centile 0-1 (les valeurs manquantes restent manquantes)."""
    return s.rank(pct=True, ascending=croissant)


def _volatilite_et_repli(cloture: pd.DataFrame, date) -> pd.DataFrame:
    """Volatilité annualisée (1 an) et pire repli sur 3 ans, par titre."""
    an = cloture.loc[:date].tail(252)
    rend = np.log(an / an.shift(1)).replace([np.inf, -np.inf], np.nan)
    vol = rend.std() * np.sqrt(252)
    trois = cloture.loc[:date].tail(756)
    repli = (trois / trois.cummax() - 1).min()
    return pd.DataFrame({"volatilite": vol, "repli_max_3a": repli})


def _per_historiques(comptes: pd.DataFrame, cloture: pd.DataFrame) -> pd.DataFrame:
    """PER de fin d'exercice : cours ajusté du dernier jour de l'année / BNPA de l'exercice."""
    fin_annee = cloture.groupby(cloture.index.year).last()
    lignes = []
    for _, r in comptes.iterrows():
        px = fin_annee[r["ticker"]].get(r["exercice"], np.nan) if r["ticker"] in fin_annee else np.nan
        per = px / r["bnpa"] if pd.notna(px) and r["bnpa"] > 0 else np.nan
        if pd.notna(per) and not PER_PLAUSIBLE[0] <= per <= PER_PLAUSIBLE[1]:
            per = np.nan
        lignes.append(per)
    return comptes.assign(per_fin_exercice=lignes)


def _note_soutenabilite(dist: float, regularite: float, rn: pd.Series, a_un_dividende: bool):
    """Note de 0 à 10 et justification. NaN si la société ne verse pas de dividende."""
    if not a_un_dividende:
        return np.nan, ["pas de dividende sur le dernier exercice"]
    pts, raisons = 0.0, []
    if pd.isna(dist):
        raisons.append("taux de distribution non calculable (résultat négatif ou absent) : 0/4")
    elif dist <= 0.60:
        pts += 4; raisons.append(f"distribution {dist:.0%} du bénéfice, marge de sécurité confortable : 4/4")
    elif dist <= 0.80:
        pts += 3; raisons.append(f"distribution {dist:.0%} du bénéfice : 3/4")
    elif dist <= 1.00:
        pts += 1.5; raisons.append(f"distribution {dist:.0%} du bénéfice, peu de réserve : 1,5/4")
    else:
        raisons.append(f"distribution {dist:.0%} : dividende supérieur au bénéfice : 0/4")
    reg = 0 if pd.isna(regularite) else regularite
    pts += 3 * reg
    raisons.append(f"dividende versé {reg:.0%} des années suivies : {3 * reg:.1f}/3".replace(".", ","))
    rn = rn.dropna()
    if len(rn) >= 3 and rn.iloc[-1] >= rn.iloc[-3]:
        pts += 2; raisons.append("résultat net stable ou en hausse sur 3 exercices : 2/2")
    elif len(rn) >= 3:
        raisons.append("résultat net en baisse sur 3 exercices : 0/2")
    if len(rn) and (rn > 0).all():
        pts += 1; raisons.append("bénéficiaire chaque exercice publié : 1/1")
    else:
        raisons.append("au moins un exercice déficitaire : 0/1")
    return round(pts, 1), raisons


def _solidite(marges: pd.Series, rn: pd.Series, tcam_ca: float):
    """Solidité économique (indicateur quantitatif de l'avantage concurrentiel)."""
    m = marges.dropna()
    pts, raisons = 0, []
    moy = m.mean() if len(m) else np.nan
    if pd.notna(moy) and moy >= 0.15:
        pts += 2; raisons.append(f"marge nette moyenne élevée ({moy:.0%})")
    elif pd.notna(moy) and moy >= 0.08:
        pts += 1; raisons.append(f"marge nette moyenne correcte ({moy:.0%})")
    elif pd.notna(moy):
        raisons.append(f"marge nette moyenne faible ({moy:.0%})")
    if len(rn.dropna()) and (rn.dropna() > 0).all():
        pts += 1; raisons.append("bénéficiaire chaque année")
    if pd.notna(tcam_ca) and tcam_ca >= 0.05:
        pts += 1; raisons.append(f"chiffre d'affaires en croissance ({tcam_ca:+.0%} par an)")
    if len(m) >= 3 and m.std() <= 0.05:
        pts += 1; raisons.append("marge stable d'une année sur l'autre")
    note = "fort" if pts >= 4 else "modéré" if pts >= 2 else "faible"
    return note, raisons


def analyser(prep: dict, res: dict, cfg: dict, saisis: pd.DataFrame | None = None) -> pd.DataFrame:
    """Une ligne par titre coté disposant de comptes publiés."""
    comptes = charger_comptes(cfg)
    mat = prep["matrices"]
    cloture = mat["cloture"]
    date = res["date"]
    tab = res["tableau"]
    comptes = _per_historiques(comptes[comptes["ticker"].isin(tab.index)], cloture)
    vr = _volatilite_et_repli(cloture, date)
    mm50 = cloture.loc[:date].tail(50).mean()
    stop = cfg["signaux"]["stop_suiveur"]

    lignes = []
    for t, g in comptes.groupby("ticker"):
        r = tab.loc[t]
        g = g.sort_values("exercice")
        dern = g.iloc[-1]
        cours = r["cours"]
        bnpa = dern["bnpa"]
        per = cours / bnpa if bnpa > 0 else np.nan
        if pd.notna(per) and not PER_PLAUSIBLE[0] <= per <= PER_PLAUSIBLE[1]:
            per = np.nan
        ca, rn = g.set_index("exercice")["chiffre_affaires_mfcfa"], g.set_index("exercice")["resultat_net_mfcfa"]
        marges = rn / ca.replace(0, np.nan)
        tcam_ca = _tcam(ca)
        dpa = dern["dividende"]
        dist = dpa / bnpa if pd.notna(dpa) and bnpa > 0 else np.nan
        sout, raisons_sout = _note_soutenabilite(dist, r["regularite"], rn, pd.notna(dpa) and dpa > 0)
        solid, raisons_solid = _solidite(marges, rn, tcam_ca)

        # fourchette à 12 mois : variations sur 12 mois réellement constatées pour ce titre
        serie = cloture[t].loc[:date].dropna().tail(10 * 252 + 1)
        r12 = (serie / serie.shift(252) - 1).dropna()
        if len(r12) >= 120:
            q_bas, q_haut = r12.quantile(0.2), r12.quantile(0.8)
            bas, haut = cours * (1 + q_bas), cours * (1 + q_haut)
            annees = len(serie) / 252
            methode = (f"sur les {annees:.0f} dernières années, le cours a fait moins bien que "
                       f"{q_bas:+.0%} sur 12 mois dans 20 % des cas et mieux que {q_haut:+.0%} "
                       "dans 20 % des cas").replace(".", ",")
        else:
            bas = haut = np.nan
            methode = "historique de cotation trop court"
        # valorisation : PER actuel face aux PER de fin d'exercice des années publiées
        per_h = g["per_fin_exercice"].dropna()
        per_moyen = per_h.mean() if len(per_h) >= 2 else np.nan

        sigma_m = vr["volatilite"].get(t, np.nan) / np.sqrt(12)
        entree_bas = cours * (1 - sigma_m) if pd.notna(sigma_m) else np.nan
        if pd.notna(mm50.get(t)) and mm50[t] < cours:
            entree_bas = max(entree_bas, mm50[t]) if pd.notna(entree_bas) else mm50[t]

        lignes.append({
            "ticker": t, "societe": r["societe"], "secteur": r["secteur"], "action": r["action"],
            "score": r["score"], "cours": cours, "exercice": int(dern["exercice"]),
            "libelle_ca": dern["libelle_ca"], "ca_dernier_mfcfa": ca.iloc[-1],
            "ca_serie": ca.round(0).tolist(), "ca_annees": [int(a) for a in ca.index],
            "tcam_ca": tcam_ca, "tcam_rn": _tcam(rn), "rn_dernier_mfcfa": rn.iloc[-1],
            "marge_nette": marges.iloc[-1], "bnpa": bnpa, "per": per,
            "rendement_12m": r["rendement_12m"], "regularite": r["regularite"],
            "dpa": dpa, "distribution": dist, "soutenabilite": sout,
            "raisons_soutenabilite": raisons_sout, "solidite": solid,
            "raisons_solidite": raisons_solid, "objectif_bas": bas, "objectif_haut": haut,
            "methode_objectif": methode, "per_moyen_hist": per_moyen, "momentum_6m": r["momentum_6m"],
            "volatilite": vr["volatilite"].get(t, np.nan),
            "repli_max_3a": vr["repli_max_3a"].get(t, np.nan),
            "montant_median_fcfa": r["montant_median_fcfa"],
            "entree_bas": entree_bas, "entree_haut": cours, "stop": cours * (1 - stop),
            "flottant_pct": dern["flottant_pct"], "detenu": r["detenu"],
            "comptes_anciens": int(dern["exercice"]) < date.year - 1,
        })
    df = pd.DataFrame(lignes).set_index("ticker")
    if df.empty:
        return df

    # PER comparé à la médiane du secteur
    df["nb_pairs"] = df.groupby("secteur")["per"].transform("count")
    df["per_secteur"] = df.groupby("secteur")["per"].transform("median").where(df["nb_pairs"] >= 3)
    df["per_vs_secteur"] = df["per"] / df["per_secteur"] - 1

    # dette / capitaux propres : uniquement si saisi (non publié par la source)
    df["dette_cp"] = np.nan
    if saisis is not None and len(saisis):
        s = saisis.copy()
        s["exercice"] = pd.to_numeric(s["exercice"], errors="coerce")
        s = s.sort_values("exercice").groupby("ticker").last()
        ratio = pd.to_numeric(s["dettes_financieres"], errors="coerce") / \
            pd.to_numeric(s["capitaux_propres"], errors="coerce")
        df["dette_cp"] = ratio.reindex(df.index)
    df["banque"] = df["secteur"].eq("Services Financiers")

    # note de risque sur 10 : rangs centiles dans l'univers analysé
    fragil = pd.Series(0.0, index=df.index)
    fragil[df["rn_dernier_mfcfa"] <= 0] = 1.0
    fragil[(df["rn_dernier_mfcfa"] > 0) & (df["tcam_rn"] < 0)] = 0.5
    fragil[df["distribution"] > 1] = np.maximum(fragil[df["distribution"] > 1], 0.5)
    comp = pd.DataFrame({
        "vol": _rang(df["volatilite"]),
        "repli": _rang(-df["repli_max_3a"]),
        "illiq": _rang(df["montant_median_fcfa"], croissant=False),
        "fragil": fragil,
    })
    poids = {"vol": 0.35, "repli": 0.25, "illiq": 0.20, "fragil": 0.20}
    df["risque"] = (sum(comp[k].fillna(0.5) * w for k, w in poids.items()) * 10).round(1)
    df["raisons_risque"] = [
        [f"volatilité {v:.0%} par an (plus agitée que {cv:.0%} des titres)" if pd.notna(v) else "volatilité non mesurable",
         f"pire repli sur 3 ans {p:.0%}" if pd.notna(p) else "repli non mesurable",
         f"montant médian échangé {m / 1e6:.1f} M FCFA par séance".replace(".", ",") if pd.notna(m) else "liquidité non mesurable",
         {1.0: "dernier exercice déficitaire", 0.5: "bénéfice en recul ou dividende non couvert",
          0.0: "bénéfices solides"}[f]]
        for v, cv, p, m, f in zip(df["volatilite"], comp["vol"].fillna(0), df["repli_max_3a"],
                                  df["montant_median_fcfa"], fragil)]
    df["potentiel_bas"] = df["objectif_bas"] / df["cours"] - 1
    df["potentiel_haut"] = df["objectif_haut"] / df["cours"] - 1
    return df


def selection_profil(df: pd.DataFrame, profil: dict, cfg: dict, n: int = 10) -> pd.DataFrame:
    """Top n adapté au profil : filtres (secteurs, risque, liquidité), puis score pondéré."""
    if df.empty:
        return df
    tol = profil["tolerance"]
    d = df[~df["action"].isin(["ÉCARTÉ"])].copy()
    if profil.get("secteurs"):
        d = d[d["secteur"].isin(profil["secteurs"])]
    d = d[d["risque"] <= RISQUE_MAX[tol]]
    if tol == "prudent":
        d = d[d["soutenabilite"].fillna(0) >= 6]
    # une ligne d'un dixième du montant doit pouvoir se vendre en quelques séances
    ligne = profil["montant_fcfa"] / n
    l = cfg["liquidite"]
    plafond = d["montant_median_fcfa"] * l["part_max_volume_journalier"] * l["seances_construction"]
    d = d[plafond >= ligne]
    composantes = {
        "score": d["score"].fillna(0) / 100,
        "securite": (10 - d["risque"]) / 10,
        "soutenabilite": d["soutenabilite"].fillna(0) / 10,
        "croissance": _rang(d["tcam_ca"]).fillna(0),
        "momentum": _rang(d["momentum_6m"]).fillna(0),
    }
    poids = PONDERATIONS_PROFIL[tol]
    d["score_profil"] = (sum(composantes[k] * w for k, w in poids.items())
                         / sum(poids.values()) * 100).round(1)
    return d.sort_values("score_profil", ascending=False).head(n)


def synthese_excel(top: pd.DataFrame, profil: dict, date) -> bytes:
    """Tableau de synthèse au format Excel (une ligne par valeur)."""
    import io

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols = {
        "societe": "Société", "secteur": "Secteur", "action": "Signal", "cours": "Cours",
        "per": "PER", "per_secteur": "PER médian secteur", "tcam_ca": "Croissance CA / an",
        "rendement_12m": "Rendement dividende", "soutenabilite": "Soutenabilité /10",
        "solidite": "Solidité", "objectif_bas": "Fourchette 12 m bas",
        "objectif_haut": "Fourchette 12 m haut", "risque": "Risque /10",
        "entree_bas": "Zone d'entrée bas", "entree_haut": "Zone d'entrée haut",
        "stop": "Stop", "score_profil": "Score profil",
    }
    t = top[list(cols)].rename(columns=cols).rename_axis("Titre").reset_index()
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        t.to_excel(w, sheet_name="Synthèse", index=False, startrow=3)
        ws = w.sheets["Synthèse"]
        ws["A1"] = f"Rapport de recherche BRVM au {date:%d/%m/%Y}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = (f"Profil : {profil['tolerance']}, {profil['montant_fcfa']:,.0f} FCFA, horizon "
                    f"{profil['horizon_mois']} mois. Outil d'aide à la décision, pas un conseil "
                    "en investissement.").replace(",", " ")
        entete = PatternFill("solid", fgColor="2B3A8C")
        for c in ws[4]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = entete
            c.alignment = Alignment(wrap_text=True, vertical="center")
        formats = {"Cours": "#,##0", "PER": "0.0", "PER médian secteur": "0.0",
                   "Croissance CA / an": "0.0%", "Rendement dividende": "0.0%",
                   "Fourchette 12 m bas": "#,##0", "Fourchette 12 m haut": "#,##0",
                   "Zone d'entrée bas": "#,##0", "Zone d'entrée haut": "#,##0", "Stop": "#,##0"}
        for j, nom in enumerate(t.columns, start=1):
            ws.column_dimensions[get_column_letter(j)].width = max(11, min(28, len(nom) + 2))
            if nom in formats:
                for i in range(5, 5 + len(t)):
                    ws.cell(i, j).number_format = formats[nom]
        ws.freeze_panes = "B5"
    return buf.getvalue()
